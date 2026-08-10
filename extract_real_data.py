# -*- coding: utf-8 -*-
"""Trich xuat ten Khoa/GV/mon hoc THAT tu TAI_NGUYEN de lam nhan cho bo sinh du lieu demo."""
import json
import re
import openpyxl
import xlrd

TAI_NGUYEN = r"D:\Cổng đào tạo\TKB\TAI_NGUYEN"

result = {"faculties": [], "teacherNames": [], "courseNames": [], "roomNames": []}

# --- 1. Ten Khoa (sheet names trong file dang ky, tru cac sheet meta) ---
wb = openpyxl.load_workbook(f"{TAI_NGUYEN}/252_Thời khóa biểu học kỳ 2 năm học 2025 - 2026.xlsx",
                            read_only=True, data_only=True)
meta_sheets = {"Hướng dẫn", "DS các học phần tại VJU", "Mẫu"}
faculties = [s for s in wb.sheetnames if s not in meta_sheets]
faculties = ["PDT" if f == "DS mon chung PDT" else f for f in faculties]
result["faculties"] = faculties
wb.close()

# --- 2. Ten mon hoc that (sheet DS cac hoc phan tai VJU, cot index 1 = "Ten") ---
wb2 = openpyxl.load_workbook(f"{TAI_NGUYEN}/252_Thời khóa biểu học kỳ 2 năm học 2025 - 2026.xlsx",
                             read_only=True, data_only=True)
sh = wb2["DS các học phần tại VJU"]
course_names = set()
for row in sh.iter_rows(min_row=2, values_only=True):
    name = row[1] if len(row) > 1 else None
    if isinstance(name, str) and name.strip():
        course_names.add(name.strip())
wb2.close()
result["courseNames"] = sorted(course_names)

# --- 3. Ten GV that + ten phong that (tu file lich-toan-truong that) ---
xls_wb = xlrd.open_workbook(f"{TAI_NGUYEN}/lich-toan-truong-14-05-2026-205950_EF0011_LichToanTruong.xls")
xsh = xls_wb.sheet_by_index(0)
teachers = set()
rooms = set()
for r in range(11, xsh.nrows):
    vals = [xsh.cell_value(r, c) for c in range(xsh.ncols)]
    if not any(str(v).strip() for v in vals):
        continue
    teacher_val = vals[11] if len(vals) > 11 else ""
    room_val = vals[10] if len(vals) > 10 else ""
    if teacher_val and isinstance(teacher_val, str):
        for name in teacher_val.split(","):
            name = name.strip()
            if name:
                teachers.add(name)
    if room_val and isinstance(room_val, str):
        for r_name in room_val.split(","):
            r_name = r_name.strip()
            if r_name:
                rooms.add(r_name)

result["teacherNames"] = sorted(teachers)
result["roomNames"] = sorted(rooms)

with open(r"D:\Cổng đào tạo\TKB\DEMO_CP_SAT\webapp\real_data.json", "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=2)

with open(r"D:\Cổng đào tạo\TKB\DEMO_CP_SAT\extract_summary.txt", "w", encoding="utf-8") as f:
    f.write(f"Faculties ({len(result['faculties'])}): {result['faculties']}\n\n")
    f.write(f"Teacher names: {len(result['teacherNames'])}\n")
    f.write(f"Sample teachers: {result['teacherNames'][:10]}\n\n")
    f.write(f"Room names: {len(result['roomNames'])}\n")
    f.write(f"Sample rooms: {result['roomNames'][:10]}\n\n")
    f.write(f"Course names: {len(result['courseNames'])}\n")
    f.write(f"Sample courses: {result['courseNames'][:15]}\n")

print("DONE")
